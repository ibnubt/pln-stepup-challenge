#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export raw tap (ts >= START_DATE) ke Excel .xlsx — konek LANGSUNG ke Postgres.
Dipakai bila server TIDAK punya Python: dijalankan di dalam container `python`
yang di-attach ke jaringan Docker proyek (lihat perintah `docker run` di catatan).

Konek pakai ENV (ada default):
  PGHOST=db  PGPORT=5432  PGUSER=postgres  PGPASSWORD=wellness  PGDATABASE=wellness
Rentang:
  START_DATE=2026-07-23   END_DATE=(kosong=sampai data terakhir)   OUT_DIR=others
"""
import os
import sys
import datetime

START_DATE = os.getenv("START_DATE", "2026-07-23")
END_DATE = os.getenv("END_DATE") or None
OUT_DIR = os.getenv("OUT_DIR", "others")
MAX_ROWS_PER_SHEET = 1_000_000


def main():
    try:
        import psycopg2
    except ImportError:
        sys.exit("psycopg2 belum ada -> pip install psycopg2-binary")
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        sys.exit("openpyxl belum ada -> pip install openpyxl")

    conn = psycopg2.connect(
        host=os.getenv("PGHOST", "db"),
        port=int(os.getenv("PGPORT", "5432")),
        user=os.getenv("PGUSER", "postgres"),
        password=os.getenv("PGPASSWORD", "wellness"),
        dbname=os.getenv("PGDATABASE", "wellness"),
    )
    where = "t.ts >= %s"
    params = [START_DATE]
    if END_DATE:
        where += " AND t.ts < %s"
        params.append(END_DATE)
    sql = (
        "SELECT t.ts, e.name, e.unit, t.level, t.kind, t.device "
        "FROM taps t JOIN employees e ON e.id = t.employee_id "
        f"WHERE {where} ORDER BY t.ts"
    )

    print(f">> query (ts >= {START_DATE}{', < ' + END_DATE if END_DATE else ''}) ...")
    cur = conn.cursor(name="stream")  # server-side cursor: streaming, hemat memori
    cur.itersize = 20000
    cur.execute(sql, params)

    header = ["ts", "name", "unit", "level", "kind", "device"]
    wb = Workbook(write_only=True)
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="12395B")

    def new_sheet(idx):
        ws = wb.create_sheet("Data" if idx == 1 else f"Data_{idx}")
        cells = []
        for h in header:
            c = WriteOnlyCell(ws, value=h)
            c.font = hfont
            c.fill = hfill
            cells.append(c)
        ws.append(cells)
        ws.freeze_panes = "A2"
        return ws

    ws = new_sheet(1)
    n = 0
    sheet_rows = 0
    sheet_idx = 1
    for row in cur:  # psycopg2 kembalikan ts sbg datetime asli
        if sheet_rows >= MAX_ROWS_PER_SHEET:
            sheet_idx += 1
            ws = new_sheet(sheet_idx)
            sheet_rows = 0
        ws.append(list(row))
        n += 1
        sheet_rows += 1
    cur.close()
    conn.close()

    os.makedirs(OUT_DIR, exist_ok=True)  # 'others/' sudah gitignore (isi PII)
    out = os.path.join(OUT_DIR, f"raw_taps_{START_DATE}_sd_{datetime.date.today().isoformat()}.xlsx")
    print(f">> menulis {n:,} baris -> {out} ...")
    wb.save(out)
    print(f">> SELESAI: {os.path.abspath(out)}  ({n:,} baris, {sheet_idx} sheet)")


if __name__ == "__main__":
    main()
