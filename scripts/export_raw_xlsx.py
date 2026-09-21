#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export SEMUA raw data tap (ts >= START_DATE s/d data terakhir) dari DB Postgres
(container Docker service `db`) ke file Excel .xlsx.

Cara pakai (di server, folder repo yang ada docker-compose.yml):
    python scripts/export_raw_xlsx.py

Prasyarat:
    - Docker jalan + stack (db) menyala.
    - Python 3 + openpyxl   ->   pip install openpyxl

Ubah rentang/filter di bagian "setelan" di bawah bila perlu.
"""
import csv
import io
import os
import sys
import subprocess
from datetime import datetime, date

# ---------------- setelan ----------------
START_DATE = "2026-07-23"        # ambil data MULAI tanggal ini (inklusif)
END_DATE = None                  # None = sampai data terakhir; atau "2026-09-22" (eksklusif)
DB_SERVICE = "db"                # nama service di docker-compose
DB_NAME = "wellness"
DB_USER = "postgres"
MAX_ROWS_PER_SHEET = 1_000_000   # limit aman Excel (maks 1.048.576/sheet)
# -----------------------------------------


def build_query():
    where = f"t.ts >= '{START_DATE}'"
    if END_DATE:
        where += f" AND t.ts < '{END_DATE}'"
    return (
        "COPY (SELECT t.ts, e.name, e.unit, t.level, t.kind, t.device "
        "FROM taps t JOIN employees e ON e.id = t.employee_id "
        f"WHERE {where} ORDER BY t.ts) TO STDOUT WITH CSV HEADER"
    )


def main():
    # pindah ke root repo (parent dari scripts/) agar `docker compose` menemukan compose file
    os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    if not (os.path.exists("docker-compose.yml") or os.path.exists("compose.yaml")):
        sys.exit("ERROR: docker-compose.yml tak ditemukan. Jalankan dari folder repo.")

    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        sys.exit("openpyxl belum terpasang. Jalankan:  pip install openpyxl")

    print(f">> menarik data (ts >= {START_DATE}{', < ' + END_DATE if END_DATE else ''}) dari DB...")
    res = subprocess.run(
        ["docker", "compose", "exec", "-T", DB_SERVICE,
         "psql", "-U", DB_USER, "-d", DB_NAME, "-c", build_query()],
        capture_output=True, encoding="utf-8",
    )
    if res.returncode != 0:
        sys.exit(f"GAGAL query DB:\n{res.stderr or res.stdout}")

    reader = csv.reader(io.StringIO(res.stdout))
    header = next(reader, None)
    if not header:
        sys.exit("Tidak ada data dikembalikan.")

    wb = Workbook(write_only=True)
    hfont = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="12395B")

    def new_sheet(idx):
        ws = wb.create_sheet("Data" if idx == 1 else f"Data_{idx}")
        cells = []
        for h in header:
            c = WriteOnlyCell(ws, value=h)
            c.font = hfont
            c.fill = hfill
            cells.append(c)
        ws.append(cells)
        ws.freeze_panes = "A2"
        return ws

    ws = new_sheet(1)
    n = 0
    sheet_rows = 0
    sheet_idx = 1
    for row in reader:
        if not row:
            continue
        # kolom 0 = ts -> datetime asli (biar Excel bisa sortir/filter tanggal)
        try:
            row[0] = datetime.strptime(row[0], "%Y-%m-%d %H:%M:%S")
        except (ValueError, IndexError):
            pass
        if sheet_rows >= MAX_ROWS_PER_SHEET:
            sheet_idx += 1
            ws = new_sheet(sheet_idx)
            sheet_rows = 0
        ws.append(row)
        n += 1
        sheet_rows += 1

    os.makedirs("others", exist_ok=True)  # folder ini sudah gitignore (isi PII)
    out = os.path.join("others", f"raw_taps_{START_DATE}_sd_{date.today().isoformat()}.xlsx")
    print(f">> menulis {n:,} baris ke {out} ...")
    wb.save(out)
    print(f">> SELESAI: {os.path.abspath(out)}  ({n:,} baris, {sheet_idx} sheet)")


if __name__ == "__main__":
    main()
